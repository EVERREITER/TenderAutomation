"""Provide shared synthetic data and fixtures for offline extraction tests.

Pytest discovers the config and sample fixtures for test configuration and a
small Excel workbook. The autouse no_network fixture blocks socket connections.
address(), extraction(), and no_missing() construct reusable contract examples.
All sample files are created in temporary test directories, without Azure calls.
"""
import socket
import pytest
from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.styles import PatternFill
from tender_extraction.config import Config
from tender_extraction.schemas import Address, SourceReference, Question, AnswerField, Extraction, Completeness


@pytest.fixture(autouse=True)
def no_network(monkeypatch):
    def blocked(*args,**kwargs): raise AssertionError("Offline-Tests dürfen kein Netzwerk verwenden")
    monkeypatch.setattr(socket.socket,"connect",blocked)
    monkeypatch.setattr(socket,"create_connection",blocked)


@pytest.fixture
def config():
    return Config(base_url="https://example.openai.azure.com/openai/v1/",api_key="secret-test-token",deployments={"luna":"dep-l","terra":"dep-t","sol":"dep-s"},di_endpoint="https://example.cognitiveservices.azure.com",di_key="di-secret-test",retries=0)


def address(cell="A1",**changes):
    return Address(**({"document":"original","sheet":"Fragen","cell_range":cell,"page":None,"bbox":None,"word_path":None,"form_field":None}|changes))


def extraction():
    refs = [SourceReference(id="s1",address=address(),quote="Provide pH"),SourceReference(id="ctx",address=address("C1"),quote="Product A")]
    q = Question(id="q1",original="Provide pH",normalized="Provide pH",language="en",kind="information",section=None,number=None,order=1,source_ids=["s1"],context_status="extracted",context_source_ids=["ctx"],position_ids=[],subquestion=None,expected_answer="number",confidence=None)
    f = AnswerField(id="f1",question_id="q1",label="pH",role="main_answer",semantic_type="number",control_type="excel_cell",target_status="located",address=address("B1"),order=1,source_ids=["s1"])
    return Extraction(questions=[q],answer_fields=[f],options=[],source_references=refs,positions=[],attributes=[],limitations=[])


@pytest.fixture
def sample(tmp_path):
    d = tmp_path/"sample_inputs"/"Unicode ü folder"
    d.mkdir(parents=True)
    p = d/"Fragen ä.xlsx"
    w = Workbook(); s = w.active; s.title="Fragen"
    s["A1"]="Provide pH"; s["C1"]="Product A"; s["D1"]="Product B"; s["A2"]="pH angeben"
    s["B1"].fill=PatternFill("solid",fgColor="FFFF00")
    s["A3"]="Combined"; s.merge_cells("A3:B3")
    s["A4"]="=1+1"
    for target, formula in [("B2",'"Yes,No"'),("B5","Choices"),("B6",'INDIRECT("D1:D2")')]:
        v=DataValidation(type="list",formula1=formula); s.add_data_validation(v); v.add(target)
    hidden=w.create_sheet("Options"); hidden.sheet_state="hidden"; hidden["A1"]="Exact YES"; hidden["A2"]="no"
    w.defined_names.add(DefinedName("Choices",attr_text="'Options'!$A$1:$A$2"))
    s.row_dimensions[7].hidden=True; s.column_dimensions["E"].hidden=True
    w.save(p); w.close()
    return tmp_path,p


def no_missing(): return Completeness(review_status="no_missing_found",missing_items=[],limitations=[])
