"""Test sparse Excel coverage, source evidence, usage reporting, and failure boundaries."""
import json
import httpx2
import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Protection
from tender_extraction.schemas import Attribute
from tender_extraction.readers.excel import read_excel
from tender_extraction.validation import validate,covered,enrich
from tender_extraction.services.openai_client import build_request
from tender_extraction.errors import ExtractionError
from conftest import extraction
from test_requests import mocked_service,response_body


def test_range_coverage_does_not_accept_one_cell_as_large_region():
    assert not covered("A1:Z100",["A1"])
    assert covered("A1:B1048576",["A1:A1048576","B1:B1048576"])
    assert not covered("A1:B3",["A1:A3","B1","B3"])


def test_static_rules_and_protection(sample):
    p=sample[1]; w=load_workbook(p); s=w["Fragen"]
    r=DataValidation(type="decimal",operator="between",formula1="0",formula2="14"); s.add_data_validation(r); r.add("B1")
    s["B1"].protection=Protection(locked=True); w.save(p); w.close()
    e=extraction(); e.attributes=[Attribute(owner_type="field",owner_id="f1",name="maximum",value=14,source_ids=["s1"])]
    checks=validate(e,read_excel(p))
    assert any(c.code=="validation_rule" and c.result=="passed" for c in checks)
    assert not any(c.code=="protected_target" for c in checks)
    e.attributes[0].value=15
    assert any(c.code=="validation_rule" and c.result=="failed" for c in validate(e,read_excel(p)))
    w=load_workbook(p); w["Fragen"].protection.sheet=True; w.save(p); w.close()
    assert any(c.code=="protected_target" for c in validate(e,read_excel(p)))


def test_invalid_json_is_not_repaired(config):
    b=response_body(); b["output"][0]["content"][0]["text"]='{"questions": ['
    s=mocked_service(config,lambda req:httpx2.Response(200,json=b))
    with pytest.raises(ExtractionError) as e: s.call("sol_extraction","excel",[])
    assert e.value.code=="schema_error" and len(s.calls)==1
    s.close()


@pytest.mark.parametrize("details,expected",[({"cached_tokens":2048},True),({"cached_tokens":0,"cache_write_tokens":2048},False),({},None)])
def test_usage_from_actual_http(config,details,expected):
    b=response_body(); b["usage"]["input_tokens_details"]=details
    s=mocked_service(config,lambda req:httpx2.Response(200,json=b))
    s.call("sol_extraction","excel",[])
    assert s.calls[0]["cache_hit"] is expected
    assert s.calls[0]["cache_write_tokens"]==details.get("cache_write_tokens")
    s.close()


def test_input_limit_prevents_request(config):
    config.max_input_bytes=1
    with pytest.raises(ExtractionError) as e: build_request(config,"sol_extraction","excel",[])
    assert e.value.code=="input_too_large"


def test_key_collision_reported_without_deleting_question(sample):
    e=extraction(); e.questions.append(e.questions[0].model_copy(update={"id":"qduplicate"})); checks=[]
    result=enrich(e,read_excel(sample[1]),{"family":"same.xlsx"},{},checks)
    assert len(result.questions)==2
    assert any(c.code=="key_collision" for c in checks)
