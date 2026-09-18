"""Regression cases for stored-cell XML streaming, independent of tender layouts."""
from datetime import datetime
from io import BytesIO
from zipfile import ZipFile
import pytest
from xml.parsers.expat import ExpatError
from openpyxl import Workbook
from openpyxl.worksheet._read_only import ReadOnlyWorksheet
from tender_extraction.readers.excel import read_excel, model_manifest, _sheet_elements, _value, NS


def replace_sheet(tmp_path, xml):
    base = tmp_path/'base.xlsx'
    workbook = Workbook()
    workbook.save(base)
    workbook.close()
    result = tmp_path/'sparse.xlsx'
    with ZipFile(base) as source, ZipFile(result, 'w') as target:
        for part in source.infolist():
            target.writestr(part, xml if part.filename=='xl/worksheets/sheet1.xml' else source.read(part.filename))
    return result


def test_no_rectangular_scan_even_with_full_excel_dimensions(tmp_path, monkeypatch):
    def reject(*args, **kwargs):
        raise AssertionError('Worksheet row iteration is forbidden for source extraction')
    monkeypatch.setattr(ReadOnlyWorksheet, 'iter_rows', reject)
    xml = f'''<worksheet xmlns="{NS['s']}">
      <dimension ref="A1:XFD1048576"/><sheetData>
      <row r="1"><c r="A1"><v>0</v></c><c r="B1" t="b"><v>0</v></c>
      <c r="C1" t="inlineStr"><is><r><t xml:space="preserve"> left </t></r><r><t>right</t></r></is></c>
      <c r="D1"><f t="shared" si="0" ref="D1:D2">A1+2</f><v>2</v></c></row>
      <row r="2"><c r="D2"><f t="shared" si="0"/><v>3</v></c><c r="G2"/></row>
      <row r="800000" hidden="1"/>
      <row r="1048576"><c r="XFD1048576" t="inlineStr"><is><t>Late content</t></is></c></row>
      </sheetData><mergeCells count="1"><mergeCell ref="G2:H2"/></mergeCells>
      </worksheet>'''.encode()
    manifest = read_excel(replace_sheet(tmp_path, xml))
    sheet = manifest['sheets'][0]
    assert len(sheet['cells']) == 6
    assert sheet['cells']['A1']['value'] == 0
    assert sheet['cells']['B1']['value'] is False
    assert sheet['cells']['C1']['value'] == ' left right'
    assert sheet['cells']['XFD1048576']['value'] == 'Late content'
    assert sheet['cells']['D1']['cached_value'] == 2
    assert sheet['cells']['D2']['formula']['attributes'] == {'t':'shared','si':'0'}
    assert sheet['merges'] == [{'range':'G2:H2','anchor':'G2'}]
    compact = model_manifest(manifest)['sheets'][0]
    assert compact['cells']['D2'] == 3
    assert compact['formulas']['D2']['attributes']['si'] == '0'
    assert compact['row_ranges'][0]['start'] == 800000
    assert 'G2' not in compact['cells']


def test_prefixed_cell_tags_and_shared_strings():
    xml = f'''<x:worksheet xmlns:x="{NS['s']}"><x:sheetData><x:row r="10">
      <x:c r="J10" t="s"><x:v>1</x:v></x:c>
      <x:c r="K10"/><x:c r="L10" t="e"><x:v>#DIV/0!</x:v></x:c>
      </x:row></x:sheetData></x:worksheet>'''.encode()
    blank, rows = {}, {}
    cells = [(el.get('r'), _value(el, ['unused','Größe & Maße'])[0])
             for el in _sheet_elements(BytesIO(xml), blank, rows) if el.tag.endswith('}c')]
    assert cells == [('J10','Größe & Maße'), ('L10','#DIV/0!')]
    assert blank[('K',0)] == [[10,10]]


def test_date_percent_and_custom_formats_survive_compaction(tmp_path):
    path = tmp_path/'formats.xlsx'
    w = Workbook()
    s = w.active
    s['J20'] = datetime(2026,9,17)
    s['A5'] = 0.25
    s['A5'].number_format = '0%'
    s['Z3'] = 1.25
    s['Z3'].number_format = '0.0000 "mg"'
    w.save(path)
    w.close()
    compact = model_manifest(read_excel(path))
    cells = compact['sheets'][0]
    assert isinstance(cells['cells']['J20'], (int,float))
    assert 'yy' in cells['cell_formats']['J20']
    assert cells['cells']['A5'] == 0.25 and cells['cell_formats']['A5'] == '0%'
    assert cells['cell_formats']['Z3'] == '0.0000 "mg"'
    assert 'date_epoch' in compact
    assert 'styles' not in compact


def test_malformed_xml_is_not_silently_treated_as_complete(tmp_path):
    xml = f'<worksheet xmlns="{NS["s"]}"><sheetData><row r="1"><c r="A1"><v>1</v></c></row>'.encode()
    path = replace_sheet(tmp_path, xml)
    with ZipFile(path) as workbook, workbook.open('xl/worksheets/sheet1.xml') as stream:
        with pytest.raises(ExpatError):
            list(_sheet_elements(stream, {}, {}))


def test_parallel_sheet_reading_preserves_complete_manifest(sample):
    serial = read_excel(sample[1], workers=1)
    parallel = read_excel(sample[1], workers=2)
    assert parallel == serial
    assert [s['name'] for s in parallel['sheets']] == ['Fragen', 'Options']
    assert parallel['sheets'][1]['state'] == 'hidden'
    assert parallel['sheets'][0]['validations'][1]['options'] == ['Exact YES', 'no']
