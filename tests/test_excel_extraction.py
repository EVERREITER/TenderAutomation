"""Test a single workbook-wide Sol request and layout-independent notes."""
import json
import httpx2
import pytest
from openpyxl import load_workbook
from tender_extraction.pipeline import run_file, digest
from tender_extraction.readers.excel import read_excel
from tender_extraction.schemas import SourceReference
from tender_extraction.validation import validate, enrich
from tender_extraction.config import Config
from conftest import extraction, address
from test_requests import mocked_service, response_body


def test_one_sol_request_contains_every_sheet_and_uses_high_budget(sample, config):
    root, path = sample
    workbook = load_workbook(path)
    workbook.copy_worksheet(workbook['Fragen']).title = 'Other product'
    workbook.save(path)
    workbook.close()
    before = digest(path)
    config.deployments = {'sol': 'dep-s'}
    candidate = extraction()
    second_q = candidate.questions[0].model_copy(deep=True, update={'id':'q2','source_ids':['s2']})
    candidate.questions.append(second_q)
    candidate.source_references.append(SourceReference(id='s2', quote='Provide pH', address=address('A1', sheet='Other product')))
    candidate.answer_fields.append(candidate.answer_fields[0].model_copy(deep=True, update={
        'id':'f2','question_id':'q2','address':address('B1', sheet='Other product')}))
    bodies = []
    def handle(request):
        bodies.append(json.loads(request.content))
        return httpx2.Response(200, json=response_body(candidate))
    service = mocked_service(config, handle)
    folder, run = run_file(path, 'outputs', config, root, service)
    service.close()
    assert run['status'] != 'failed', run.get('error')
    assert digest(path) == before
    assert len(bodies) == 1
    assert bodies[0]['model'] == 'dep-s'
    assert bodies[0]['reasoning'] == {'effort':'high'}
    assert bodies[0]['max_output_tokens'] == 125000
    payload = json.loads(bodies[0]['input'][1]['content'][0]['text'])
    assert 'target_sheet' not in payload
    sheets = payload['source_manifest']['sheets']
    assert [s['name'] for s in sheets] == ['Fragen', 'Options', 'Other product']
    assert sheets[1]['state'] == 'hidden'
    assert sheets[0]['validations'][1]['options'] == ['Exact YES', 'no']
    assert {p.name for p in folder.iterdir()} == {'sol_questions.json', 'result.json'}
    result = json.loads((folder/'result.json').read_text(encoding='utf-8'))
    assert [q['id'] for q in result['questions']] == ['q1','q2']
    assert result['questions'][1]['answer_fields'][0]['address']['sheet'] == 'Other product'
    assert run['stages'] == ['sol_extraction','python_validation']
    assert run['excel_source']['sheets'] == 3


def test_workbook_truncation_does_not_save_partial_extraction(sample, config):
    service = mocked_service(config, lambda req: httpx2.Response(200, json=response_body(
        status='incomplete', incomplete_details={'reason':'max_output_tokens'})))
    folder, run = run_file(sample[1], 'outputs', config, sample[0], service)
    service.close()
    assert run['error']['code'] == 'truncation'
    assert run['status'] == 'failed'
    assert len(run['calls']) == 1
    assert json.loads((folder/'result.json').read_text())['questions'] == []
    assert not (folder/'sol_questions.json').exists()


def test_completed_but_empty_response_with_limitations_is_failed(sample, config):
    from tender_extraction.schemas import Extraction, Position
    candidate = Extraction(questions=[], answer_fields=[], options=[], source_references=[],
        positions=[Position(id='p1', product='Example', lot=None, strength=None, pack=None,
                            form=None, source_ids=['missing_source'])], attributes=[],
        limitations=['Recognized requirements omitted because the answer might be too long.'])
    service = mocked_service(config, lambda req: httpx2.Response(200, json=response_body(candidate)))
    folder, run = run_file(sample[1], 'outputs', config, sample[0], service)
    service.close()
    assert run['status'] == 'failed'
    assert run['error']['code'] == 'empty_extraction'
    assert len(run['calls']) == 1
    assert run['calls'][0]['response_status'] == 'completed'
    assert any(c['code']=='reference_integrity' for c in run['validation_checks'])
    assert json.loads((folder/'sol_questions.json').read_text())['limitations'] == candidate.limitations
    assert json.loads((folder/'result.json').read_text())['questions'] == []


def test_genuinely_empty_candidate_without_limitations_is_allowed(sample, config):
    from tender_extraction.schemas import Extraction
    from openpyxl import Workbook
    workbook = Workbook()
    workbook.save(sample[1])
    workbook.close()
    candidate = Extraction(questions=[], answer_fields=[], options=[], source_references=[],
                           positions=[], attributes=[], limitations=[])
    service = mocked_service(config, lambda req: httpx2.Response(200, json=response_body(candidate)))
    _, run = run_file(sample[1], 'outputs', config, sample[0], service)
    service.close()
    assert run['status'] == 'completed'


@pytest.mark.parametrize('note_cell,note',[
    ('Z40', 'Immer aufrunden. Konzentrationsgrenzen angeben.'),
    ('A100', 'Specify temperature and attach the certificate.'),
    ('K2', 'Indiquez la concentration et la durée de conservation.'),
])
def test_notes_are_sourced_by_model_without_fixed_headers_or_columns(sample, config, note_cell, note):
    root, path = sample
    workbook = load_workbook(path)
    workbook['Fragen'][note_cell] = note
    workbook.save(path)
    workbook.close()
    candidate = extraction()
    candidate.questions[0].notes = note
    candidate.questions[0].note_source_ids = ['note']
    candidate.source_references.append(SourceReference(id='note', quote=note, address=address(note_cell)))
    service = mocked_service(config, lambda req: httpx2.Response(200, json=response_body(candidate)))
    folder, run = run_file(path, 'outputs', config, root, service)
    service.close()
    assert run['status'] != 'failed', run.get('error')
    result = json.loads((folder/'result.json').read_text(encoding='utf-8'))
    assert result['questions'][0]['notes'] == note
    assert any(c['code']=='notes_evidence' and c['result']=='passed' for c in result['questions'][0]['checks'])
    assert not any(c['code']=='requirement_coverage' for c in run['validation_checks'])


def test_invented_notes_are_flagged(sample):
    candidate = extraction()
    candidate.questions[0].notes = 'Invented instruction'
    candidate.questions[0].note_source_ids = ['missing']
    checks = validate(candidate, read_excel(sample[1]))
    assert any(c.code=='notes_evidence' and c.result=='failed' for c in checks)


def test_sheet_name_remains_a_valid_context_source(sample):
    candidate = extraction()
    candidate.source_references[1].address = address(None)
    candidate.source_references[1].quote = 'Fragen'
    manifest = read_excel(sample[1])
    checks = validate(candidate, manifest)
    assert any(c.code=='sheet_name_quote' and c.result=='passed' for c in checks)
    assert enrich(candidate, manifest, {'family':'test.xlsx'}, {}, checks).schema_version == '3'


def test_output_budget_can_be_overridden_without_reading_project_secrets(monkeypatch, tmp_path):
    monkeypatch.chdir(tmp_path)
    monkeypatch.delenv('EXTRACTION_MAX_OUTPUT_TOKENS', raising=False)
    assert Config.from_env().extraction_tokens == 125000
    monkeypatch.setenv('EXTRACTION_MAX_OUTPUT_TOKENS', '64000')
    assert Config.from_env().extraction_tokens == 64000


def test_prompt_example_matches_schema_and_source_evidence(tmp_path):
    from importlib.resources import files
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill
    from tender_extraction.schemas import Extraction
    prompt = files('tender_extraction').joinpath('prompts/excel_workbook.md').read_text(encoding='utf-8')
    candidate = Extraction.model_validate_json(prompt.split('```json\n', 1)[1].split('```', 1)[0])
    path = tmp_path / 'example.xlsx'
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Angebot'
    # Reconstruct exactly the source stated by the example, not model guesses.
    for ref in candidate.source_references:
        sheet[ref.address.cell_range] = ref.quote
    sheet['C3'].fill = PatternFill('solid', fgColor='FFFF00')
    workbook.save(path)
    workbook.close()
    checks = validate(candidate, read_excel(path))
    assert all(c.result == 'passed' for c in checks), [c.model_dump() for c in checks]
    assert candidate.questions[0].position_ids == ['p1']
    assert candidate.questions[0].notes == candidate.source_references[1].quote
