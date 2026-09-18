"""Test Excel content preservation, dropdowns, product contexts, and source validation."""
import copy
from zipfile import ZipFile
import pytest
from tender_extraction.readers.excel import read_excel, model_manifest, contains
from tender_extraction.validation import validate,enrich
from tender_extraction.schemas import Option
from conftest import extraction,address


def test_structure_and_lists(sample):
    root,path=sample; m=read_excel(path); s=m["sheets"][0]
    assert len(m["sheets"])==2 and m["sheets"][1]["state"]=="hidden"
    assert [r["options"] for r in s["validations"]]==[["Yes","No"],["Exact YES","no"],None]
    assert s["cells"]["A4"]["formula"]["text"]=="1+1" and s["cells"]["A4"]["cached_value"] is None
    assert s["merges"]==[{"range":"A3:B3","anchor":"A3"}]
    assert s["row_ranges"] and s["columns"] and s["blank_ranges"]


def test_formatting_tail_is_not_sent_but_distant_content_survives(tmp_path):
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill
    from openpyxl.comments import Comment
    from openpyxl.worksheet.datavalidation import DataValidation
    w = Workbook()
    s = w.active
    s["A1"] = "Question"
    s["B1"].fill = PatternFill("solid", fgColor="FFFF00")
    s["B1048576"].fill = copy.copy(s["B1"].fill)
    s["A900000"] = "Distant question"
    s["A900001"] = 0
    s["A900002"] = False
    s["A900003"] = "=1+1"
    s["C700000"].comment = Comment("Separate instruction", "Author")
    s["A5"] = "  keep spaces  "
    rule = DataValidation(type="list", formula1='"Yes,No"')
    s.add_data_validation(rule)
    rule.add("D800000")
    p = tmp_path / "tail.xlsx"
    w.save(p)
    full = read_excel(p)
    compact = model_manifest(full)
    local, sent = full["sheets"][0], compact["sheets"][0]
    assert "B1048576:B1048576" in [r["range"] for r in local["blank_ranges"]]
    assert "blank_ranges" not in sent
    assert sent["cells"].keys() == local["cells"].keys()
    for address, cell in local["cells"].items():
        assert sent["cells"][address] == (cell["cached_value"] if cell["formula"] is not None else cell["value"])
        assert sent["formulas"].get(address) == cell["formula"]
    assert sent["cells"]["A5"] == "  keep spaces  "
    assert sent["comments"] == local["comments"]
    assert sent["validations"] == local["validations"]
    assert "raw_structure" not in sent


def test_multiple_fields_products_and_parallel_language(sample):
    m=read_excel(sample[1]); e=extraction()
    e.source_references.append(e.source_references[0].model_copy(update={"id":"de","quote":"pH angeben","address":address("A2")}))
    e.questions[0].source_ids.append("de")
    e.answer_fields.append(e.answer_fields[0].model_copy(update={"id":"comment","role":"comment","address":address("B2")}))
    e.source_references.append(e.source_references[1].model_copy(update={"id":"ctxB","quote":"Product B","address":address("D1")}))
    e.questions.append(e.questions[0].model_copy(update={"id":"q2","context_source_ids":["ctxB"]}))
    e.answer_fields.append(e.answer_fields[0].model_copy(update={"id":"f2","question_id":"q2","address":address("B5")}))
    c=validate(e,m); r=enrich(e,m,{"family":"test.xlsx"},{},c)
    assert len(r.questions)==2 and len(r.questions[0].answer_fields)==2
    assert r.questions[0].key!=r.questions[1].key
    e.questions.reverse()
    r2=enrich(e,m,{"family":"test.xlsx"},{},[])
    assert {q.id:q.key for q in r.questions}=={q.id:q.key for q in r2.questions}


@pytest.mark.parametrize("mutation,code,result",[("missing","missing_target","not_verifiable"),("unknown","unknown_context","not_verifiable"),("invalid","excel_address","failed"),("merge","merge_anchor","failed"),("fake_dropdown","dropdown_type","failed"),("dynamic","dropdown_options","not_verifiable"),("reference","reference_integrity","failed"),("formula","formula_target","failed")])
def test_validation_cases(sample,mutation,code,result):
    m=read_excel(sample[1]); e=extraction()
    if mutation=="missing": e.answer_fields=[]
    if mutation=="unknown": e.questions[0].context_status="unknown"
    if mutation=="invalid": e.source_references[0].address.sheet="invented"
    if mutation=="merge": e.answer_fields[0].address.cell_range="B3"
    if mutation=="fake_dropdown": e.answer_fields[0].control_type="excel_dropdown"
    if mutation=="dynamic": e.answer_fields[0].address.cell_range="B6"
    if mutation=="reference": e.questions[0].source_ids=["missing"]
    if mutation=="formula": e.answer_fields[0].address.cell_range="A4"
    checks=validate(e,m)
    assert any(c.code==code and c.result==result for c in checks)


def test_exact_options_and_shared_target(sample):
    e=extraction(); e.answer_fields[0].address=address("B2"); e.answer_fields[0].control_type="excel_dropdown"
    e.options=[Option(id=str(i),field_id="f1",label=v,value=v,meaning=None,order=i,source_ids=["s1"]) for i,v in enumerate(["Yes","No"])]
    m=read_excel(sample[1]); assert any(c.code=="dropdown_options" and c.result=="passed" for c in validate(e,m))
    e.options.reverse(); e.options[0].order=0; e.options[1].order=1
    assert any(c.code=="dropdown_options" and c.result=="not_verifiable" for c in validate(e,m))
    e.answer_fields.append(e.answer_fields[0].model_copy(update={"id":"f2"}))
    assert any(c.code=="shared_physical_target" for c in validate(e,m))
